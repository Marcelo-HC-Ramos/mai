"""UC03: relatório de inteligência de mercado em HTML (imprimível em PDF) e XLSX."""
from __future__ import annotations

import io
from datetime import timezone

from jinja2 import Environment, PackageLoader, select_autoescape

from .analysis.recommendations import brl
from .connectors.base import SHIPPING_MODES

SOURCE_LABEL = {"api": "Dado do marketplace", "calculado": "Calculado pelo MAI", "estimativa": "Estimativa",
                "vendedor": "Dados do vendedor", "indisponivel": "Indisponível"}

_env = Environment(loader=PackageLoader("mai", "templates"), autoescape=select_autoescape(["html"]))
_env.filters["brl"] = brl
_env.filters["src"] = lambda s: SOURCE_LABEL.get(s, s)
_env.filters["pct"] = lambda v: "-" if v is None else f"{v:.1f}%".replace(".", ",")


def price_histogram_svg(prices: list[float], markers: dict[str, float | None], width: int = 640, height: int = 180) -> str:
    """Histograma dos preços dos concorrentes com linhas para o preço atual e o recomendado."""
    if not prices:
        return ""
    lo, hi = min(prices), max(prices)
    vals = [v for v in markers.values() if v]
    lo, hi = min([lo] + vals), max([hi] + vals)
    span = (hi - lo) or 1
    bins = 12
    counts = [0] * bins
    for p in prices:
        counts[min(bins - 1, int((p - lo) / span * bins))] += 1
    top = max(counts) or 1
    pad_l, pad_b, pad_t = 8, 28, 26
    plot_w, plot_h = width - 2 * pad_l, height - pad_b - pad_t
    bw = plot_w / bins
    parts = [f'<svg viewBox="0 0 {width} {height}" role="img" aria-label="Distribuição de preços dos concorrentes" '
             f'class="hist">']
    for i, c in enumerate(counts):
        h = plot_h * c / top
        x = pad_l + i * bw + 1
        y = pad_t + plot_h - h
        label = f"{brl(lo + span * i / bins)} a {brl(lo + span * (i + 1) / bins)}: {c} anúncio(s)"
        if c:
            parts.append(f'<rect x="{x:.1f}" y="{y:.1f}" width="{bw - 2:.1f}" height="{h:.1f}" rx="3" '
                         f'class="bar"><title>{label}</title></rect>')
    base_y = pad_t + plot_h
    parts.append(f'<line x1="{pad_l}" x2="{width - pad_l}" y1="{base_y}" y2="{base_y}" class="axis"/>')
    parts.append(f'<text x="{pad_l}" y="{height - 8}" class="tick">{brl(lo)}</text>')
    parts.append(f'<text x="{width - pad_l}" y="{height - 8}" class="tick" text-anchor="end">{brl(hi)}</text>')
    styles = {"Seu preço": "m-current", "Recomendado": "m-reco", "Mediana": "m-median"}
    for name, value in markers.items():
        if not value:
            continue
        x = pad_l + plot_w * (value - lo) / span
        parts.append(f'<line x1="{x:.1f}" x2="{x:.1f}" y1="{pad_t - 6}" y2="{base_y}" '
                     f'class="marker {styles.get(name, "")}"><title>{name}: {brl(value)}</title></line>')
    parts.append("</svg>")
    return "".join(parts)


def render_html(snapshot: dict, product: dict) -> str:
    comp = snapshot.get("competitor_data", {})
    pricing = snapshot.get("pricing", {})
    hist = price_histogram_svg(comp.get("price_distribution", []), {
        "Seu preço": pricing.get("current_price"), "Recomendado": pricing.get("recommended_price"),
        "Mediana": comp.get("median_price")})
    captured = snapshot["captured_at"]
    if captured.tzinfo is None:
        captured = captured.replace(tzinfo=timezone.utc)
    return _env.get_template("report.html").render(
        s=snapshot, p=product, comp=comp, pricing=pricing,
        kw=snapshot["search_intelligence"]["keywords"], dev=snapshot["search_intelligence"]["traffic_breakdown"],
        rev=snapshot["review_insights"], recs=snapshot["ai_recommendations"], hist=hist, modes=SHIPPING_MODES,
        captured=captured.strftime("%d/%m/%Y %H:%M UTC"))


def render_xlsx(snapshot: dict, product: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    bold = Font(bold=True)

    def sheet(title, header, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font = bold
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(70, max(12, *(len(str(c.value or "")) for c in col)))
        return ws

    comp, pricing = snapshot["competitor_data"], snapshot["pricing"]
    kw = snapshot["search_intelligence"]["keywords"]
    dev = snapshot["search_intelligence"]["traffic_breakdown"]
    sheet("Resumo", ["Indicador", "Valor", "Fonte"], [
        ["Produto", product["title"], ""],
        ["Data da análise", snapshot["captured_at"].strftime("%d/%m/%Y %H:%M"), ""],
        ["Seu preço", pricing.get("current_price"), "api"],
        ["Preço recomendado", pricing.get("recommended_price"), "calculado"],
        ["Preço competitivo", pricing.get("competitive_price"), "calculado"],
        ["Preço mínimo permitido", pricing.get("min_allowed_price"), "calculado"],
        ["Margem no recomendado (%)", pricing.get("margin_pct_at_recommended"), "calculado"],
        ["Concorrentes analisados", comp.get("competitors_used"), "api"],
        ["Menor preço", comp.get("lowest_price"), "api"],
        ["Mediana", comp.get("median_price"), "calculado"],
        ["Média", comp.get("average_price"), "calculado"],
        ["Maior preço", comp.get("highest_price"), "api"],
        ["Buscas pelo celular (%)", dev["mobile_share_pct"], dev["source"]],
        ["Buscas pelo computador (%)", dev["desktop_share_pct"], dev["source"]],
    ], first=True)
    sheet("Concorrentes", ["Anúncio", "Título", "Preço", "Vendidos", "Envio", "Frete grátis", "Reputação", "Link"],
          [[c["item_id"], c["title"], c["price"], c["sold_quantity"], SHIPPING_MODES.get(c.get("shipping_mode"), "-"),
            "Sim" if c["free_shipping"] else "Não", c["seller_reputation"], c["url"]] for c in comp.get("top_competitors", [])])
    logi = snapshot.get("logistics") or {}
    if logi.get("by_mode"):
        sheet("Full x Flex", ["Modalidade", "Anúncios", "% dos anúncios", "% das vendas", "Preço mediano", "Vendidos por anúncio"],
              [[m["label"], m["listings"], m["share_pct"], m["sales_share_pct"], m["median_price"], m["avg_sold"]]
               for m in logi["by_mode"]])
    breakdown = (pricing.get("cost_breakdown") or {}).get("breakdown")
    if breakdown:
        sheet("Custos", ["Item", "Valor (R$)", "% do preço", "Fonte"],
              [[b["label"], b["value"], b["pct_of_price"], b["source"]] for b in breakdown]
              + [["Lucro líquido", pricing["cost_breakdown"]["profit"], pricing["cost_breakdown"]["margin_pct"], ""]])
    sheet("Palavras-chave", ["Termo", "Pontuação (0-100)", "Volume mensal", "Já no seu título", "Fontes"],
          [[k["term"], k["score"], k["monthly_volume"], "Sim" if k["in_my_title"] else "Não", ", ".join(k["sources"])]
           for k in kw["top_keywords"]])
    rows = []
    for owner, g in snapshot["review_insights"]["groups"].items():
        for r in g["reasons"]:
            rows.append([owner, r["label"], r["count"], r["pct_of_bad_reviews"], " | ".join(r["examples"])])
    sheet("Avaliações ruins", ["Grupo", "Motivo", "Ocorrências", "% das ruins", "Exemplos"], rows)
    plan = [[r["priority"], r["type"], r["message"]] for r in snapshot["ai_recommendations"]]
    ai = snapshot.get("ai_summary")
    if ai:
        plan = [[0, "RESUMO (IA)", ai["summary"]], [0, "TÍTULO (IA)", ai.get("title_suggestion")]] + \
               [[0, "ARGUMENTO (IA)", sp] for sp in ai.get("selling_points", [])] + plan
    sheet("Plano de ação", ["Prioridade", "Tipo", "Recomendação"], plan)
    sheet("Curva de preço", ["Preço", "Conversão relativa", "Margem unitária", "Lucro esperado", "Permitido"],
          [[r["price"], r["conversion"], r["unit_margin"], r["expected_profit"], "Sim" if r["allowed"] else "Não"]
           for r in pricing.get("curve", [])])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
